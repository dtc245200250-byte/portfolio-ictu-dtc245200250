import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def create_report():
    doc = docx.Document()

    # 1. Page Margins (A4 standard: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 2.0 cm
        section.bottom_margin = Inches(0.79)  # 2.0 cm
        section.left_margin = Inches(1.18)    # 3.0 cm
        section.right_margin = Inches(0.79)   # 2.0 cm
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

    # 2. Base Normal Style Font Setup
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 41, 59) # #1E293B

    # Helper styling functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138) # #1E3A8A (Navy Blue)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 235) # #2563EB (Royal Blue)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
        return p

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
            r_bold.font.bold = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.italic = italic
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(12)
            r_bold.font.bold = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(code_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_image_figure(img_path, caption_text, width_inches=5.8):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=Inches(width_inches))

            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            run = p_cap.add_run(caption_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        else:
            p_err = doc.add_paragraph(f"[Ảnh minh họa: {caption_text} - Đường dẫn: {img_path}]")
            p_err.runs[0].font.italic = True

    # ==============================================================================
    # 1. TRANG BÌA (OFFICIAL ICTU COVER PAGE)
    # ==============================================================================
    p_header1 = doc.add_paragraph()
    p_header1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header1.paragraph_format.space_after = Pt(2)
    r1 = p_header1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — ĐẠI HỌC THÁI NGUYÊN")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True

    p_header2 = doc.add_paragraph()
    p_header2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header2.paragraph_format.space_after = Pt(18)
    r2 = p_header2.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(30, 58, 138)

    # Official ICTU Logo
    logo_path = r"c:\All\Triển khai_Thi\docs\assets\logo_ictu_full.png"
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(20)
        p_logo.add_run().add_picture(logo_path, width=Inches(3.2))

    p_title1 = doc.add_paragraph()
    p_title1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title1.paragraph_format.space_before = Pt(10)
    p_title1.paragraph_format.space_after = Pt(8)
    r_t1 = p_title1.add_run("BÁO CÁO BÀI TẬP LỚN\nTRIỂN KHAI VÀ QUẢN TRỊ HỆ THỐNG PHẦN MỀM")
    r_t1.font.name = 'Times New Roman'
    r_t1.font.size = Pt(18)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(30, 58, 138)

    p_title2 = doc.add_paragraph()
    p_title2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title2.paragraph_format.space_after = Pt(24)
    r_t2 = p_title2.add_run("ĐỀ TÀI 13:\nXÂY DỰNG WEBSITE PORTFOLIO CÁ NHÂN VÀ\nTRIỂN KHAI HỆ THỐNG BẰNG DOCKER")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(16)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(220, 38, 38) # Crimson accent

    # Info Box Table
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    
    col_widths = [Inches(2.2), Inches(3.6)]
    labels = [
        "Sinh viên thực hiện:",
        "Mã số sinh viên (MSSV):",
        "Lớp / Khóa:",
        "Ngành học:",
        "Học phần:"
    ]
    vals = [
        "Phan Quang Tùng",
        "DTC245200250",
        "CNTTK23E / Khóa K23",
        "Công nghệ Thông tin",
        "Triển khai và Quản trị Hệ thống Phần mềm"
    ]

    for idx in range(5):
        row = info_table.rows[idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(labels[idx])
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(12)
        r0.font.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        r1 = p1.add_run(vals[idx])
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        if idx == 0 or idx == 1:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(30, 58, 138)

    # Footer date
    p_footer = doc.add_paragraph()
    p_footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_before = Pt(40)
    r_f = p_footer.add_run("Thái Nguyên, Năm 2026")
    r_f.font.name = 'Times New Roman'
    r_f.font.size = Pt(12)
    r_f.font.bold = True

    doc.add_page_break()

    # ==============================================================================
    # 2. LỜI MỞ ĐẦU & DANH MỤC VIẾT TẮT
    # ==============================================================================
    add_heading_1("LỜI MỞ ĐẦU")
    add_body("Trong xu thế chuyển đổi số và công nghiệp phần mềm hiện đại, việc xây dựng và phát triển một ứng dụng không chỉ dừng lại ở mã nguồn (source code) hay giao diện đồ họa người dùng. Yếu tố cốt lõi quyết định khả năng vận hành ổn định, khả năng mở rộng (scalability), tính sẵn sàng cao (high availability) và mức độ an toàn thông tin nằm ở năng lực triển khai (Deployment) và quản trị hạ tầng (System Administration).")
    add_body("Trước đây, việc triển khai ứng dụng lên các máy chủ vật lý hoặc máy chủ ảo truyền thống (Virtual Machines) thường đối mặt với vấn đề nan giải: 'chạy được trên máy lập trình viên nhưng lỗi trên máy chủ sản xuất' (It works on my machine problem), sự xung đột phiên bản thư viện, môi trường hệ điều hành không đồng nhất và lãng phí tài nguyên tính toán. Công nghệ Container hóa (Containerization) với đầu tàu là Docker ra đời đã tạo nên một cuộc cách mạng trong văn hóa phát triển phần mềm và DevOps, cho phép đóng gói ứng dụng cùng toàn bộ runtime, thư viện phụ thuộc và cấu hình vào một đơn vị độc lập, khả chuyển và nhất quán tuyệt đối.")
    add_body("Báo cáo bài tập lớn với Đề tài số 13: 'Website Portfolio Cá nhân và Triển khai Hệ thống bằng Docker' được sinh viên Phan Quang Tùng (MSSV: DTC245200250, Lớp CNTTK23E) nghiên cứu, thiết kế và trực tiếp triển khai trong môn học 'Triển khai và Quản trị Hệ thống Phần mềm'. Dự án không chỉ là một trang web đơn thuần mà là một hạ tầng Microservices hoàn chỉnh gồm 13 dịch vụ container liên kết chặt chẽ: từ cổng Reverse Proxy Nginx bảo mật SSL/TLS, ứng dụng Portfolio full-stack (React + Node.js/Express + MySQL), công cụ quản trị dữ liệu trực quan phpMyAdmin, cho đến hệ thống giám sát hiệu năng thời gian thực (Prometheus, Grafana, cAdvisor, Node Exporter, MySQL Exporter, Nginx Exporter) và trung tâm thu thập, phân tích log tập trung (Loki, Promtail, LogQL).")
    add_body("Sinh viên xin bày tỏ lòng cảm ơn chân thành đến các Thầy, Cô giáo Bộ môn Hệ thống Thông tin & Mạng máy tính, Khoa Công nghệ Thông tin, Trường Đại học Công nghệ Thông tin và Truyền thông đã tận tình truyền đạt kiến thức nền tảng và phương pháp thực nghiệm quý báu trong suốt học kỳ.")

    add_heading_2("BẢNG DANH MỤC THUẬT NGỮ VÀ VIẾT TẮT")
    
    term_table = doc.add_table(rows=1, cols=3)
    term_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    term_table.autofit = False
    
    headers = ["Thuật ngữ", "Tên đầy đủ / Ý nghĩa", "Vai trò trong Đề tài 13"]
    hdr_cells = term_table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    terms_data = [
        ("CI/CD", "Continuous Integration / Continuous Deployment", "Tự động hóa tích hợp và phát hành phần mềm."),
        ("CRUD", "Create, Read, Update, Delete", "Các thao tác quản lý dữ liệu Portfolio qua Admin Panel."),
        ("cAdvisor", "Container Advisor (Google)", "Công cụ thu thập tài nguyên CPU/RAM/IO của từng Container."),
        ("Exporter", "Prometheus Metrics Exporter", "Tiến trình thu thập metrics chuyên biệt (Node, MySQL, Nginx)."),
        ("JWT", "JSON Web Token", "Cơ chế xác thực không trạng thái (stateless) cho Admin API."),
        ("LogQL", "Log Query Language (Grafana Loki)", "Ngôn ngữ truy vấn và lọc dữ liệu log phân tán."),
        ("Loki", "Grafana Loki Log Aggregation System", "Hệ thống lưu trữ log tập trung tối ưu hóa metadata index."),
        ("Prometheus", "Prometheus Time-Series Database", "Hệ thống cào và lưu trữ số liệu hiệu năng (Metrics TSDB)."),
        ("Promtail", "Grafana Promtail Log Shipper", "Agent thu thập log từ Docker engine và đẩy về Loki."),
        ("Reverse Proxy", "Nginx HTTP/HTTPS Reverse Proxy", "Cổng Gateway tiếp nhận request, cân bằng tải, bảo mật SSL."),
        ("SAN", "Subject Alternative Name", "Mở rộng chứng chỉ SSL cho phép bảo vệ tên miền/IP cục bộ.")
    ]

    for row_idx, (t, f, r) in enumerate(terms_data):
        row = term_table.add_row()
        cells = row.cells
        bg = "FFFFFF" if row_idx % 2 == 0 else "F1F5F9"
        for i, val in enumerate([t, f, r]):
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], top=80, bottom=80, left=100, right=100)
            p = cells[i].paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if i == 0:
                run.font.bold = True

    doc.add_page_break()

    # ==============================================================================
    # 3. CHƯƠNG 1
    # ==============================================================================
    add_heading_1("CHƯƠNG 1: TỔNG QUAN KIẾN TRÚC HỆ THỐNG VÀ CHUẨN BỊ MÔI TRƯỜNG")
    
    add_heading_2("1.1. Bối cảnh và Mục tiêu đề tài")
    add_body("Đề tài 13 tập trung xây dựng một hệ thống Website Portfolio Cá nhân hiện đại dành cho sinh viên CNTT, nhằm trình diễn hồ sơ năng lực, các dự án phần mềm, kỹ năng chuyên môn, kinh nghiệm thực tập và thông tin liên hệ. Thay vì chỉ dựng giao diện tĩnh thông thường, đề tài đặt ra yêu cầu kỹ thuật toàn diện theo mô hình chuẩn sản xuất (Production-Ready) của môn học:")
    add_bullet("Đóng gói độc lập toàn bộ các tầng ứng dụng (Frontend, Backend, Database) bằng Docker Container.", "Containerization: ");
    add_bullet("Cấu hình Nginx làm cổng tiếp nhận duy nhất (Reverse Proxy), kích hoạt SSL/TLS tự ký và gia cố bảo mật bằng HTTP Security Headers.", "Reverse Proxy & SSL: ");
    add_bullet("Xây dựng hệ thống giám sát Metrics thời gian thực với Prometheus và trực quan hóa qua Grafana Dashboard.", "Observability (Metrics): ");
    add_bullet("Triển khai hệ thống lưu vết và phân tích Log tập trung với Loki và Promtail, kiểm chứng qua các truy vấn LogQL.", "Centralized Logging: ");
    add_bullet("Áp dụng các tiêu chuẩn an toàn thông tin: Non-root user, cô lập mạng nội bộ (Network Isolation), bảo vệ biến môi trường bí mật qua file .env.", "Hardening & Security: ");
    add_bullet("Quản lý vòng đời mã nguồn trên GitHub theo chuẩn Git commit milestone phản ánh trung thực quá trình thực hiện.", "Source Code Management: ");

    add_heading_2("1.2. Sơ đồ Kiến trúc Tổng thể Hệ thống")
    add_body("Hệ sinh thái được thiết kế theo kiến trúc Microservices phân lớp rõ ràng, bao gồm tổng cộng 13 Docker Containers hoạt động đồng bộ:")
    
    add_body("Dưới đây là bảng phân bổ chi tiết 13 Container Services trong hệ thống:")

    srv_table = doc.add_table(rows=1, cols=4)
    srv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    srv_table.autofit = False
    
    s_headers = ["Tên Container", "Công nghệ / Image", "Cổng nội bộ / Public", "Vai trò và Chức năng"]
    s_hdr_cells = srv_table.rows[0].cells
    for i, h in enumerate(s_headers):
        set_cell_background(s_hdr_cells[i], "1E3A8A")
        set_cell_margins(s_hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = s_hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    srv_data = [
        ("portfolio_nginx", "nginx:1.27-alpine", "8080:80, 8443:443", "Cổng Reverse Proxy, SSL Offloading, Security Headers"),
        ("portfolio_frontend", "trinkhai_thi-frontend", "80 (Internal)", "Giao diện Portfolio SPA React 18, Vite, TypeScript"),
        ("portfolio_backend", "trinkhai_thi-backend", "5000 (Internal)", "RESTful API Node.js, Express, TypeScript, JWT Auth"),
        ("portfolio_mysql", "mysql:8.0", "3306 (Internal)", "Hệ quản trị CSDL quan hệ lưu trữ dữ liệu Portfolio"),
        ("portfolio_phpmyadmin", "phpmyadmin:5.2.2", "8088:80", "Giao diện Web trực quan quản lý CSDL MySQL"),
        ("portfolio_prometheus", "prom/prometheus:v2.54.1", "9090:9090", "Máy chủ lưu trữ và cào số liệu hiệu năng (Metrics TSDB)"),
        ("portfolio_grafana", "grafana/grafana:11.2.0", "3000:3000", "Bảng điều khiển trực quan hóa Dashboard giám sát"),
        ("portfolio_cadvisor", "gcr.io/cadvisor:v0.49.2", "8080 (Internal)", "Thu thập số liệu CPU, RAM, Network I/O từng container"),
        ("portfolio_node_exporter", "prom/node-exporter:v1.8.2", "9100 (Internal)", "Thu thập số liệu tài nguyên phần cứng máy chủ host"),
        ("portfolio_mysqld_exporter", "prom/mysqld-exporter:v0.15.1", "9104 (Internal)", "Thu thập hiệu năng truy vấn, connections của MySQL"),
        ("portfolio_nginx_exporter", "nginx/nginx-prometheus-exporter", "9113 (Internal)", "Thu thập active connections, request rate từ Nginx"),
        ("portfolio_loki", "grafana/loki:3.0.0", "3100:3100", "Máy chủ lưu trữ và lập chỉ mục Log tập trung"),
        ("portfolio_promtail", "grafana/promtail:3.0.0", "Daemon (Host)", "Agent thu thập log từ Docker engine socket gửi về Loki")
    ]

    for row_idx, row_data in enumerate(srv_data):
        row = srv_table.add_row()
        cells = row.cells
        bg = "FFFFFF" if row_idx % 2 == 0 else "F1F5F9"
        for i, val in enumerate(row_data):
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], top=70, bottom=70, left=90, right=90)
            p = cells[i].paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            if i == 0:
                run.font.bold = True

    add_heading_2("1.3. Mô hình Phân vùng Mạng Nội bộ (Docker Network Isolation)")
    add_body("Để tuân thủ triệt để tiêu chuẩn bảo mật Hardening (Bước 6), hệ thống không sử dụng một mạng bridge mặc định chung mà chia tách thành 3 phân vùng mạng Docker độc lập (Multi-network Architecture):")
    add_bullet("frontend_net: Dành riêng cho giao tiếp giữa Nginx Reverse Proxy, Frontend SPA và Backend API. Tầng người dùng ngoài chỉ có thể đi qua Nginx.", "Phân vùng Ngoại vi (Frontend): ");
    add_bullet("backend_net: Dành riêng cho kết nối giữa Backend API, MySQL Database và MySQL Exporter. Mạng này hoàn toàn cách ly với Internet và Frontend.", "Phân vùng Dữ liệu (Backend): ");
    add_bullet("monitoring_net: Kết nối các thành phần giám sát (Prometheus, Grafana, Loki, Promtail và toàn bộ các Exporters) để truyền tải số liệu và log an toàn.", "Phân vùng Giám sát (Monitoring): ");

    add_heading_2("1.4. Quản lý Mã nguồn trên GitHub theo Milestone")
    add_body("Toàn bộ mã nguồn, cấu hình Docker, kịch bản tự động hóa và tài liệu được quản lý tập trung trên GitHub tại địa chỉ: https://github.com/dtc245200250-byte/portfolio-ictu-dtc245200250")
    add_body("Quy trình Git được tổ chức theo đúng các mốc hoàn thành (Milestones) yêu cầu trong đề cương môn học:")
    add_bullet("Commit 1: feat(app): trien khai ung dung portfolio, co so du lieu mysql, phpmyadmin va nginx reverse proxy", "Mốc 1 (Bước 1 - Bước 3): ");
    add_bullet("Commit 2: feat(monitoring): tich hop he thong giam sat prometheus, grafana, cadvisor va exporters", "Mốc 2 (Bước 4): ");
    add_bullet("Commit 3: feat(logging): tich hop he thong quan ly log tap trung loki va promtail kem logql", "Mốc 3 (Bước 5 & 6): ");
    add_bullet("Commit 4: docs: bo sung bao cao word de tai 13, screenshots, assets va huong dan readme day du", "Mốc Hoàn thiện Tài liệu: ");

    doc.add_page_break()

    # ==============================================================================
    # 4. CHƯƠNG 2
    # ==============================================================================
    add_heading_1("CHƯƠNG 2: TRIỂN KHAI ỨNG DỤNG WEB, CƠ SỞ DỮ LIỆU VÀ REVERSE PROXY")

    add_heading_2("2.1. Thiết kế và Xây dựng Ứng dụng Portfolio (Frontend SPA)")
    add_body("Ứng dụng giao diện người dùng được xây dựng theo kiến trúc Single Page Application (SPA) sử dụng React 18, Vite và TypeScript, mang lại hiệu năng tải trang vượt trội, tính an toàn kiểu dữ liệu (Type-safety) và trải nghiệm thẩm mỹ cao cấp:")
    add_bullet("Trình bày đầy đủ thông tin sinh viên Phan Quang Tùng, MSSV DTC245200250, sinh viên CNTTK23E Khoa CNTT - ICTU.", "Hero Profile: ");
    add_bullet("Phân loại các kỹ năng theo nhóm: Frontend (React, TS, Tailwind), Backend (Node.js, Express, Python), Database (MySQL, Redis), DevOps (Docker, Linux, CI/CD, Nginx).", "Chuyên môn & Kỹ năng: ");
    add_bullet("Danh sách các dự án thực tế kèm huy hiệu công nghệ (tech badges), mô tả kiến trúc và liên kết GitHub.", "Dự án Nổi bật: ");
    add_bullet("Mục tiêu nghề nghiệp trở thành Kỹ sư Phần mềm Full-stack & DevOps chuyên nghiệp.", "Mục tiêu & Định hướng: ");
    add_bullet("Hệ thống quản trị CRUD bảo mật bằng JWT, cho phép xem, thêm, sửa, xóa các dự án và kỹ năng theo thời gian thực.", "Trang Quản trị (Admin Panel): ");

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\01_portfolio_homepage.png",
        "Hình 2.1: Giao diện Trang chủ Website Portfolio Cá nhân của sinh viên Phan Quang Tùng (Cổng 8080)"
    )

    add_heading_2("2.2. Xây dựng Dịch vụ Backend RESTful API")
    add_body("Backend đóng vai trò xử lý logic nghiệp vụ và giao tiếp cơ sở dữ liệu, được viết bằng Node.js kết hợp Express framework và TypeScript. Mã nguồn được cấu trúc theo mô hình phân lớp chuẩn (Layered Architecture): Controllers, Services, Middleware, Routes và Database Helpers.")
    add_body("Các tính năng bảo mật then chốt được triển khai tại Backend API:")
    add_bullet("Sử dụng thư viện bcrypt để băm mật khẩu quản trị viên với độ phức tạp salt rounds = 10, tuyệt đối không lưu trữ mật khẩu ở dạng plain-text.", "Mã hóa Mật khẩu: ");
    add_bullet("Phát hành JSON Web Token có chữ ký bảo mật (JWT_SECRET) và thời hạn hiệu lực 24 giờ cho các tác vụ CRUD nhạy cảm.", "Xác thực Phiên JWT: ");
    add_bullet("Thư viện Helmet được tích hợp để tự động áp dụng 11 HTTP security headers tại tầng ứng dụng.", "Helmet Protection: ");
    add_bullet("Sử dụng express-rate-limit giới hạn tối đa 100 requests / 15 phút trên mỗi địa chỉ IP nhằm ngăn chặn tấn công từ chối dịch vụ (DoS) và brute-force.", "Chống Brute-force: ");
    add_bullet("Tích hợp endpoint chuyên dụng /metrics xuất số liệu theo định dạng Prometheus (process memory, heap, uptime).", "Prometheus Metrics Endpoint: ");

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\02_portfolio_admin_login.png",
        "Hình 2.2: Giao diện Đăng nhập Quản trị viên (Admin Login) bảo vệ bằng JWT Token"
    )

    add_heading_2("2.3. Thiết kế Cơ sở Dữ liệu Quan hệ MySQL 8 và phpMyAdmin")
    add_body("Hệ quản trị cơ sở dữ liệu MySQL phiên bản 8.0 được lựa chọn nhờ độ tin cậy cao, hỗ trợ giao dịch ACID toàn vẹn và tối ưu hóa tài nguyên. Cơ sở dữ liệu portfolio_db được cấu trúc gồm các bảng chính:")
    add_bullet("users: Lưu trữ tài khoản quản trị (username, password_hash, email, full_name, role).", "Bảng Quản trị viên: ");
    add_bullet("profile: Lưu trữ thông tin cá nhân của sinh viên Phan Quang Tùng (MSSV, lớp, khoa, trường, tiểu sử, thông tin liên hệ).", "Bảng Hồ sơ Cá nhân: ");
    add_bullet("skills: Lưu danh mục kỹ năng, phân loại (frontend, backend, database, devops) và mức độ thành thạo.", "Bảng Kỹ năng Chuyên môn: ");
    add_bullet("projects: Lưu thông tin các dự án thực hiện (tên dự án, mô tả, công nghệ, link github, demo, thời gian thực hiện).", "Bảng Dự án: ");
    add_bullet("contact_messages: Lưu trữ các tin nhắn liên hệ do người xem website gửi đến.", "Bảng Tin nhắn Liên hệ: ");

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\06_phpmyadmin.png",
        "Hình 2.3: Giao diện Quản trị Cơ sở Dữ liệu MySQL bằng phpMyAdmin (Cổng 8088)"
    )

    add_heading_2("2.4. Cấu hình Nginx Reverse Proxy, SSL/TLS và Security Headers")
    add_body("Nginx được cấu hình như một lá chắn phòng thủ tại cổng giao tiếp ngoại vi. Mọi kết nối từ máy khách đều đi qua Nginx trước khi chuyển tiếp vào các container nội bộ. Điều này mang lại các lợi ích vượt trội:")
    add_bullet("Ẩn giấu hoàn toàn địa chỉ IP và cấu trúc nội bộ của các container ứng dụng (Backend, Frontend, phpMyAdmin).", "Ẩn giấu Hạ tầng: ");
    add_bullet("Chứng chỉ SSL/TLS tự ký (Self-signed Certificate) được sinh với chuẩn mã hóa RSA 2048-bit và Subject Alternative Name (SAN=localhost, IP=127.0.0.1), phục vụ giao thức HTTPS an toàn trên cổng 8443.", "Mã hóa HTTPS: ");
    add_bullet("Thiết lập bộ Security Headers chuẩn OWASP phòng chống các cuộc tấn công mạng nguy hiểm:", "Gia cố HTTP Headers: ");

    add_code_block("""# Cấu hình Security Headers trên Nginx (file: nginx/conf.d/default.conf)
add_header X-Frame-Options "SAMEORIGIN" always;
add_header X-Content-Type-Options "nosniff" always;
add_header X-XSS-Protection "1; mode=block" always;
add_header Referrer-Policy "strict-origin-when-cross-origin" always;
add_header Content-Security-Policy "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; font-src 'self' https://fonts.gstatic.com data:; img-src 'self' data: https:;" always;
ssl_protocols TLSv1.2 TLSv1.3;
server_tokens off;""")

    doc.add_page_break()

    # ==============================================================================
    # 5. CHƯƠNG 3
    # ==============================================================================
    add_heading_1("CHƯƠNG 3: GIÁM SÁT HỆ THỐNG VỚI PROMETHEUS, GRAFANA VÀ EXPORTERS")

    add_heading_2("3.1. Kiến trúc Hệ thống Giám sát Metrics")
    add_body("Khả năng quan sát (Observability) là trụ cột không thể thiếu trong quản trị hệ thống phần mềm hiện đại. Kiến trúc giám sát trong Đề tài 13 hoạt động theo cơ chế 'Pull-based' chuẩn của Prometheus: Prometheus Server định kỳ gửi HTTP GET request đến các endpoint /metrics của các Exporters theo chu kỳ cào (scrape interval) 5 giây.")

    add_heading_2("3.2. Cấu hình Prometheus Server và Scrape Configs")
    add_body("File cấu hình prometheus.yml thiết lập 6 Scrape Targets độc lập, bao quát toàn diện từ tầng phần cứng, container, cơ sở dữ liệu đến ứng dụng:")

    add_code_block("""# Cấu hình Scrape Targets (file: monitoring/prometheus/prometheus.yml)
global:
  scrape_interval: 5s
  evaluation_interval: 5s

scrape_configs:
  - job_name: 'prometheus'
    static_configs:
      - targets: ['localhost:9090']

  - job_name: 'cadvisor'
    static_configs:
      - targets: ['cadvisor:8080']

  - job_name: 'node-exporter'
    static_configs:
      - targets: ['node-exporter:9100']

  - job_name: 'mysql-exporter'
    static_configs:
      - targets: ['mysqld-exporter:9104']

  - job_name: 'nginx-exporter'
    static_configs:
      - targets: ['nginx-exporter:9113']

  - job_name: 'portfolio-backend'
    static_configs:
      - targets: ['backend:5000']
    metrics_path: '/metrics'""")

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\03_prometheus_targets.png",
        "Hình 3.1: Danh sách toàn bộ 6 Scrape Targets của Prometheus đều đạt trạng thái UP (Health: 100%)"
    )

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\04_prometheus_graph.png",
        "Hình 3.2: Giao diện Prometheus Graph hiển thị biểu đồ metric 'up' của 6 thành phần"
    )

    add_heading_2("3.3. Xây dựng Bảng điều khiển Giám sát trên Grafana")
    add_body("Grafana phiên bản 11.2.0 được tích hợp kết nối với 2 Datasources chính: Prometheus (Metrics) và Loki (Logs). Thông qua cơ chế Provisioning tự động (không cần cấu hình thủ công bằng tay), Grafana tự động nạp Dashboard giám sát chuyên sâu 'Hạ tầng Container & Observability - DTC245200250':")
    add_bullet("Trạng thái Hoạt động Toàn hệ thống (Service Health Gauge): Hiển thị tỷ lệ phần trăm dịch vụ UP/DOWN.", "Health Indicator: ");
    add_bullet("Mức độ Tiêu thụ CPU theo từng Container (cAdvisor): Đồ thị đường thể hiện % CPU theo thời gian thực.", "CPU Utilization: ");
    add_bullet("Mức độ Tiêu thụ Bộ nhớ RAM (cAdvisor): Dung lượng RAM thực tế đang dùng của từng microservice.", "Memory Usage: ");
    add_bullet("Lưu lượng Mạng I/O Container: Số bytes truyền và nhận qua các card mạng ảo.", "Network Traffic: ");
    add_bullet("Tần suất Truy vấn Cơ sở dữ liệu MySQL: Tốc độ xử lý câu lệnh SELECT, INSERT, UPDATE.", "Database QPS: ");

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\05_grafana_overview.png",
        "Hình 3.3: Bảng điều khiển Grafana Dashboard giám sát trực quan tài nguyên hạ tầng 13 Container"
    )

    doc.add_page_break()

    # ==============================================================================
    # 6. CHƯƠNG 4
    # ==============================================================================
    add_heading_1("CHƯƠNG 4: HỆ THỐNG QUẢN TRỊ LOG TẬP TRUNG LOKI VÀ PROMTAIL")

    add_heading_2("4.1. Sự cần thiết của Quản lý Log tập trung")
    add_body("Trong một hệ thống bao gồm 13 container hoạt động đồng thời, việc người quản trị phải SSH vào từng container hoặc gõ lệnh docker logs cho từng dịch vụ để khắc phục sự cố (troubleshooting) là vô cùng thủ công, thiếu hiệu quả và dễ bỏ sót thông tin khi container bị khởi động lại.")
    add_body("Giải pháp kết hợp giữa Grafana Loki và Promtail được ví như 'Prometheus dành cho Logs', đem lại các ưu điểm vượt trội:")
    add_bullet("Loki chỉ lập chỉ mục (index) các nhãn metadata (như job, container, level) tương tự Prometheus, không đánh chỉ mục toàn văn (full-text index) giúp tiết kiệm đến 80% dung lượng đĩa cứng.", "Tối ưu hóa Tài nguyên: ");
    add_bullet("Tích hợp mượt mà với giao diện Grafana Explore, cho phép liên kết trực tiếp giữa biểu đồ Metrics bất thường và các dòng Log tương ứng tại đúng thời điểm đó.", "Khả năng Tương quan (Correlation): ");

    add_heading_2("4.2. Cấu hình Promtail Log Shipper và Docker Socket Discovery")
    add_body("Promtail được cấp quyền truy cập vào Docker socket (/var/run/docker.sock) và thư mục logs chia sẻ (shared_logs) để tự động phát hiện mọi container đang chạy trên host và đẩy dữ liệu log về Loki Server:")

    add_code_block("""# Cấu hình Promtail Scrape Configs (file: monitoring/promtail/promtail-config.yml)
clients:
  - url: http://loki:3100/loki/api/v1/push

scrape_configs:
  - job_name: docker-containers
    docker_sd_configs:
      - host: unix:///var/run/docker.sock
        refresh_interval: 5s
    relabel_configs:
      - source_labels: ['__meta_docker_container_name']
        regex: '/(.*)'
        target_label: 'container'
      - source_labels: ['__meta_docker_container_name']
        regex: '/(.*)'
        target_label: 'job'""")

    add_heading_2("4.3. Xác thực và Đánh giá 3 Truy vấn LogQL Tiêu biểu")
    add_body("Hệ thống đã thực hiện kiểm thử thực tế và xác thực thành công 3 câu truy vấn LogQL tiêu biểu theo đúng yêu cầu đề tài:")
    
    add_bullet("{job=\"portfolio_nginx\"} — Truy vấn toàn bộ luồng log truy cập và lỗi phát sinh từ cổng Nginx Reverse Proxy. Giúp người quản trị nắm bắt lưu lượng truy cập tổng quan.", "Truy vấn 1 (Toàn bộ Log Nginx): ");
    add_bullet("{job=\"portfolio_nginx\"} |~ \"(GET|POST)\" — Sử dụng bộ lọc biểu thức chính quy (Regex filter) để bóc tách chính xác các HTTP requests phương thức GET và POST gửi đến hệ thống, loại bỏ các dòng log khởi động không liên quan.", "Truy vấn 2 (Lọc HTTP Methods): ");
    add_bullet("{job=\"portfolio_backend\"} — Lọc dữ liệu log ghi nhận từ ứng dụng Backend API, ghi nhận chi tiết thời gian phản hồi, mã trạng thái HTTP 200/400/500 và các truy vấn cơ sở dữ liệu.", "Truy vấn 3 (Log Ứng dụng Backend): ");

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\08_loki_logql_validated.png",
        "Hình 4.1: Kết quả thực thi kiểm chứng thành công 3 truy vấn LogQL trên Loki Engine bằng script kiểm thử tự động"
    )

    doc.add_page_break()

    # ==============================================================================
    # 7. CHƯƠNG 5
    # ==============================================================================
    add_heading_1("CHƯƠNG 5: CÁC BIỆN PHÁP HARDENING, KIỂM THỬ VÀ ĐÁNH GIÁ")

    add_heading_2("5.1. Triển khai các Biện pháp Hardening Bảo mật")
    add_body("Nhằm bảo vệ hệ thống trước các mối đe dọa an ninh mạng phổ biến và đáp ứng yêu cầu của Bước 6, các giải pháp gia cố bảo mật (System Hardening) đã được áp dụng toàn diện:")
    add_bullet("Mặc định các tiến trình trong container Docker chạy dưới quyền root (UID 0), tiềm ẩn nguy cơ kẻ tấn công thoát khỏi container (container escape) để chiếm quyền máy chủ host. Dự án đã cấu hình user nginx (trong Frontend và Nginx) và cấu hình phân quyền nghiêm ngặt nhằm triệt tiêu hoàn toàn rủi ro này.", "Chạy Container với quyền Non-root User: ");
    add_bullet("Tách biệt 3 mạng nội bộ (frontend_net, backend_net, monitoring_net). Các dịch vụ nhạy cảm như MySQL Database không mở cổng (port mapping) ra ngoài máy chủ host, chỉ chấp nhận kết nối từ Backend và Exporter trong cùng mạng.", "Cô lập Mạng (Network Segmentation): ");
    add_bullet("Toàn bộ thông tin nhạy cảm (mật khẩu root MySQL, mật khẩu database user, khóa ký JWT, tài khoản Grafana) được tách riêng vào tệp .env. Tệp .env được đưa vào danh sách đen trong .gitignore để tuyệt đối không đẩy lên kho mã nguồn công khai.", "Bảo vệ Khóa Bí mật (.env & .gitignore): ");
    add_bullet("Nginx được cấu hình server_tokens off nhằm ẩn hoàn toàn phiên bản phần mềm trong HTTP headers, gây khó khăn cho kẻ tấn công khi quét lỗ hổng đã biết (CVE).", "Ẩn Phiên bản Phần mềm (Server Tokens Off): ");
    add_bullet("Sử dụng multi-stage Docker build cho cả Frontend và Backend để loại bỏ hoàn toàn mã nguồn TypeScript, trình biên dịch và các file rác khỏi image sản phẩm cuối cùng, giảm thiểu diện tích tấn công (Attack Surface).", "Multi-stage Docker Builds: ");

    add_heading_2("5.2. Kết quả Trạng thái Vận hành 13 Containers")
    add_body("Hình ảnh chụp thực tế từ terminal lệnh 'docker compose ps' chứng minh tất cả 13 containers đều đang trong trạng thái Up, các dịch vụ có healthcheck (MySQL, Backend, Frontend, cAdvisor) đều đạt trạng thái Healthy:")

    add_image_figure(
        r"c:\All\Triển khai_Thi\docs\screenshots\07_docker_compose_ps.png",
        "Hình 5.1: Bằng chứng thực tế 13 Containers Docker đều đang vận hành ổn định và Healthy trên hệ thống"
    )

    add_heading_2("5.3. Bảng Đánh giá Kiểm thử Toàn diện Hệ thống")
    add_body("Dưới đây là bảng tổng hợp kết quả kiểm thử chức năng và phi chức năng của toàn bộ hạ tầng:")

    test_table = doc.add_table(rows=1, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.autofit = False
    
    t_headers = ["Hạng mục Kiểm thử", "Phương pháp Thực hiện", "Kết quả Kỹ thuật Đạt được", "Đánh giá"]
    t_hdr_cells = test_table.rows[0].cells
    for i, h in enumerate(t_headers):
        set_cell_background(t_hdr_cells[i], "1E3A8A")
        set_cell_margins(t_hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = t_hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    test_data = [
        ("Truy cập HTTP/HTTPS", "Gửi request cổng 8080 và 8443", "Nginx trả về mã 200 OK, SSL SAN hợp lệ", "ĐẠT (PASS)"),
        ("Hiển thị Portfolio", "Tải trang chủ trên trình duyệt Edge", "Dữ liệu sinh viên Phan Quang Tùng hiển thị đầy đủ", "ĐẠT (PASS)"),
        ("Xác thực Admin JWT", "Đăng nhập với mật khẩu mã hóa", "Cấp Token 24h, mở quyền truy cập CRUD", "ĐẠT (PASS)"),
        ("Cơ sở dữ liệu MySQL", "Kiểm tra kết nối và truy vấn bảng", "Tạo lập 5 bảng, seed dữ liệu mẫu thành công", "ĐẠT (PASS)"),
        ("phpMyAdmin Web UI", "Đăng nhập cổng 8088", "Giao diện trực quan xem bảng và cấu trúc DB", "ĐẠT (PASS)"),
        ("Scrape Metrics", "Kiểm tra endpoint /targets", "Tất cả 6 targets Prometheus đều UP 100%", "ĐẠT (PASS)"),
        ("Grafana Dashboard", "Mở dashboard qua cổng 3000", "Hiển thị đầy đủ biểu đồ CPU, RAM, Network", "ĐẠT (PASS)"),
        ("Thu thập Log tập trung", "Chạy script kiểm thử LogQL", "3 truy vấn LogQL trả về dữ liệu chuẩn xác", "ĐẠT (PASS)"),
        ("Network Isolation", "Kiểm tra mạng nội bộ Docker", "MySQL không mở port ra host, chỉ nội bộ kết nối", "ĐẠT (PASS)"),
        ("Git Milestone Commits", "Kiểm tra git log trên branch main", "Tối thiểu 3 commits theo đúng tiến trình đề cương", "ĐẠT (PASS)")
    ]

    for row_idx, row_data in enumerate(test_data):
        row = test_table.add_row()
        cells = row.cells
        bg = "FFFFFF" if row_idx % 2 == 0 else "F1F5F9"
        for i, val in enumerate(row_data):
            set_cell_background(cells[i], bg)
            set_cell_margins(cells[i], top=70, bottom=70, left=90, right=90)
            p = cells[i].paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            if i == 3:
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 101, 52) # Dark Green

    doc.add_page_break()

    # ==============================================================================
    # 8. KẾT LUẬN VÀ TÀI LIỆU THAM KHẢO
    # ==============================================================================
    add_heading_1("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    add_body("Sau thời gian nghiên cứu lý thuyết, phân tích yêu cầu, trực tiếp lập trình và triển khai thực tế trên hệ thống, sinh viên Phan Quang Tùng (MSSV: DTC245200250) đã hoàn thành xuất sắc toàn bộ các mục tiêu đặt ra của Đề tài số 13 — 'Website Portfolio Cá nhân và Triển khai Hệ thống bằng Docker'.")
    
    add_heading_2("Các kết quả nổi bật đạt được:")
    add_bullet("Xây dựng hoàn chỉnh ứng dụng Portfolio full-stack hiện đại bằng React 18, Node.js/Express, TypeScript và MySQL 8 với đầy đủ các tính năng trình diễn và quản trị CRUD chuyên nghiệp.", "Về Ứng dụng: ");
    add_bullet("Thiết lập kiến trúc Microservices container hóa gồm 13 dịch vụ hoạt động ổn định, phối hợp nhịp nhàng thông qua Docker Compose và phân vùng mạng nội bộ nghiêm ngặt.", "Về Đóng gói & Triển khai: ");
    add_bullet("Hoàn thiện hệ sinh thái quan sát toàn diện (Observability) kết hợp chặt chẽ giữa số liệu hiệu năng (Prometheus Metrics) và dấu vết sự kiện (Loki Centralized Logs) trên cùng giao diện Grafana.", "Về Giám sát & Quản trị: ");
    add_bullet("Áp dụng nghiêm ngặt các nguyên tắc Hardening an toàn thông tin: Non-root user, bảo vệ bí mật tệp .env, mã hóa HTTPS tự ký với SAN và bộ HTTP Security Headers chuẩn OWASP.", "Về Bảo mật: ");

    add_heading_2("Hướng phát triển và mở rộng trong tương lai:")
    add_bullet("Nâng cấp từ Docker Compose đơn máy chủ lên cụm điều phối Docker Swarm hoặc Kubernetes (K8s) nhằm hỗ trợ tự động cân bằng tải (Load Balancing) và tự phục hồi khi có sự cố nút (Self-healing).", "Điều phối Cụm (Orchestration): ");
    add_bullet("Thiết lập đường ống tự động hóa kiểm thử và triển khai liên tục (CI/CD Pipeline) bằng GitHub Actions kết hợp công cụ quét lỗ hổng bảo mật container (Trivy, Docker Scout).", "Tự động hóa CI/CD: ");
    add_bullet("Cấu hình Prometheus Alertmanager để tự động gửi thông báo cảnh báo qua Telegram Bot hoặc email quản trị viên khi phát sinh sự cố vượt ngưỡng tài nguyên (CPU > 85%, RAM > 90%).", "Cảnh báo Tự động (Alerting): ");

    add_heading_1("TÀI LIỆU THAM KHẢO")
    refs = [
        "Trường Đại học Công nghệ Thông tin và Truyền thông (ICTU) — Đề cương môn học Triển khai và Quản trị Hệ thống Phần mềm, 2026.",
        "Docker Inc. — Official Docker Documentation & Best Practices for Production Containerization, 2026 (https://docs.docker.com).",
        "Nginx Documentation — Reverse Proxy, SSL Termination and Security Headers Configuration Guide (https://nginx.org/en/docs/).",
        "Prometheus Authors — Prometheus Monitoring System and Time Series Database Documentation (https://prometheus.io/docs/).",
        "Grafana Labs — Grafana, Loki and Promtail Observability Architecture Documentation (https://grafana.com/docs/).",
        "Open Web Application Security Project (OWASP) — REST Security Cheat Sheet & Secure Headers Project, 2026."
    ]
    for idx, r in enumerate(refs, 1):
        add_bullet(r, bold_prefix=f"[{idx}] ")

    output_file = r"c:\All\Triển khai_Thi\Bao_Cao_De_Tai_13_Phan_Quang_Tung.docx"
    doc.save(output_file)
    print(f"Report successfully generated and saved to: {output_file}")

if __name__ == "__main__":
    create_report()
